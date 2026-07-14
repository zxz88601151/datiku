# © 中哥  All Rights Reserved
# 版权标识: FP_UUID_31adb5871aea40b8b0c288773f094ab2|FP_AUTHOR_中哥_SN_20260531|FP_HASH_20260531B9F3|FP_ORIGIN_2026_AUTHOR_中哥
# 仅限项目内部使用，未经授权禁止转载、商用。

"""Excel 解析器。

依赖 openpyxl（可选）：未安装时给出明确提示而非静默失败。
Excel 列定义见 EXCEL_TEMPLATE.md，列名容错（忽略大小写/空格）。
"""

from __future__ import annotations

from typing import List, Dict, Any


def parse_excel(path: str) -> List[Dict[str, Any]]:
    """读取 .xlsx，返回归一化题目行列表。"""
    try:
        import openpyxl  # noqa
    except ImportError:
        raise RuntimeError(
            "解析 Excel 需要 openpyxl，请先安装：pip install openpyxl"
        )
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(h).strip() if h is not None else "" for h in rows[0]]

    # 列名 → 标准字段 的模糊映射
    alias = {
        "年级": "grade", "学科": "subject", "章节": "chapter", "知识点": "knowledge",
        "题型": "type", "题目": "content", "题干": "content", "答案": "answer",
        "解析": "analysis", "难度": "difficulty",
        "选项a": "opt_a", "选项b": "opt_b", "选项c": "opt_c", "选项d": "opt_d",
        "正确选项": "correct", "正确答案": "correct",
    }

    out: List[Dict[str, Any]] = []
    for raw in rows[1:]:
        if raw is None or all(c is None for c in raw):
            continue
        rec: Dict[str, Any] = {}
        for col, val in zip(header, raw):
            key = alias.get(col.strip().lower())
            if key:
                rec[key] = val
        out.append(rec)
    return out


def _build_options(rec: Dict[str, Any]) -> List[Dict[str, Any]]:
    """从 选项A..D / 正确选项 列构造 options 列表。"""
    opts = []
    correct = str(rec.get("correct", "")).strip().upper()
    for label in ("A", "B", "C", "D"):
        text = rec.get(f"opt_{label.lower()}")
        if text is None or str(text).strip() == "":
            continue
        opts.append({
            "label": label,
            "text": str(text).strip(),
            "correct": label == correct,
        })
    return opts
